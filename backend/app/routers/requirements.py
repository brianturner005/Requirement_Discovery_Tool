import csv
import io
import math
from datetime import date
from typing import List

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.dependencies import get_current_user
from app.database import get_db
from app.models.user import User
from app.schemas.common import PaginatedResponse, StatusEnum
from app.schemas.requirement import (
    RequirementCreate,
    RequirementResponse,
    RequirementSummary,
    RequirementUpdate,
    StatusTransition,
)
from app.schemas.requirement_version import VersionResponse
from app.services import requirement_service as svc
from app.services import requirement_version_service as ver_svc
from app.services import jira_service, linear_service


class BulkStatusRequest(BaseModel):
    req_ids: List[str]
    status: str

router = APIRouter(prefix="/requirements", tags=["requirements"])


def _to_response(req) -> RequirementResponse:
    related = list({r.req_id: r for r in list(req.related_to) + list(req.related_from)}.values())
    return RequirementResponse(
        id=req.id,
        req_id=req.req_id,
        title=req.title,
        description=req.description,
        source=req.source,
        stakeholder=req.stakeholder,
        system=req.system,
        priority=req.priority,
        confidence=req.confidence,
        business_impact=req.business_impact,
        technical_impact=req.technical_impact,
        status=req.status,
        notes=req.notes,
        tags=req.tags,
        related_requirements=[
            RequirementSummary(id=r.id, req_id=r.req_id, title=r.title, status=r.status, priority=r.priority)
            for r in related
        ],
        evidence=req.evidence,
        jira_issue_key=req.jira_issue_key,
        linear_issue_id=req.linear_issue_id,
        created_at=req.created_at,
        updated_at=req.updated_at,
    )


@router.get("", response_model=PaginatedResponse[RequirementResponse])
async def list_requirements(
    q: str | None = Query(None, description="Full-text search"),
    status: list[str] | None = Query(None),
    priority: list[str] | None = Query(None),
    source: list[str] | None = Query(None),
    system_id: list[int] | None = Query(None),
    stakeholder_id: list[int] | None = Query(None),
    tag: list[str] | None = Query(None),
    confidence: list[str] | None = Query(None),
    sort_by: str = Query("updated_at"),
    sort_dir: str = Query("desc"),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
):
    items, total = await svc.list_requirements(
        db,
        q=q,
        status=status,
        priority=priority,
        source=source,
        system_id=system_id,
        stakeholder_id=stakeholder_id,
        tag=tag,
        confidence=confidence,
        sort_by=sort_by,
        sort_dir=sort_dir,
        page=page,
        page_size=page_size,
    )
    return PaginatedResponse(
        items=[_to_response(r) for r in items],
        total=total,
        page=page,
        page_size=page_size,
        pages=max(1, math.ceil(total / page_size)),
    )


@router.get("/export")
async def export_requirements(
    q: str | None = Query(None),
    status: list[str] | None = Query(None),
    priority: list[str] | None = Query(None),
    source: list[str] | None = Query(None),
    system_id: list[int] | None = Query(None),
    stakeholder_id: list[int] | None = Query(None),
    tag: list[str] | None = Query(None),
    confidence: list[str] | None = Query(None),
    db: AsyncSession = Depends(get_db),
):
    items, _ = await svc.list_requirements(
        db,
        q=q, status=status, priority=priority, source=source,
        system_id=system_id, stakeholder_id=stakeholder_id,
        tag=tag, confidence=confidence,
        page=1, page_size=10_000,
    )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "ID", "Title", "Status", "Priority", "Source", "Confidence",
        "Stakeholder", "System", "Tags", "Business Impact",
        "Technical Impact", "Notes", "Created", "Updated",
    ])
    for r in items:
        writer.writerow([
            r.req_id,
            r.title,
            r.status,
            r.priority,
            r.source,
            r.confidence,
            r.stakeholder.name if r.stakeholder else "",
            r.system.name if r.system else "",
            ", ".join(t.name for t in r.tags),
            r.business_impact or "",
            r.technical_impact or "",
            r.notes or "",
            r.created_at.strftime("%Y-%m-%d"),
            r.updated_at.strftime("%Y-%m-%d"),
        ])

    output.seek(0)
    filename = f"requirements-{date.today().isoformat()}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/xlsx")
async def export_requirements_xlsx(
    q: str | None = Query(None),
    status: list[str] | None = Query(None),
    priority: list[str] | None = Query(None),
    source: list[str] | None = Query(None),
    system_id: list[int] | None = Query(None),
    stakeholder_id: list[int] | None = Query(None),
    tag: list[str] | None = Query(None),
    confidence: list[str] | None = Query(None),
    db: AsyncSession = Depends(get_db),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=503, detail="openpyxl not installed")

    items, _ = await svc.list_requirements(
        db, q=q, status=status, priority=priority, source=source,
        system_id=system_id, stakeholder_id=stakeholder_id,
        tag=tag, confidence=confidence, page=1, page_size=10_000,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Requirements"

    headers = ["ID", "Title", "Status", "Priority", "Source", "Confidence",
               "Stakeholder", "System", "Tags", "Business Impact",
               "Technical Impact", "Notes", "Created", "Updated"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for r in items:
        ws.append([
            r.req_id, r.title, r.status, r.priority, r.source, r.confidence,
            r.stakeholder.name if r.stakeholder else "",
            r.system.name if r.system else "",
            ", ".join(t.name for t in r.tags),
            r.business_impact or "", r.technical_impact or "", r.notes or "",
            r.created_at.strftime("%Y-%m-%d"), r.updated_at.strftime("%Y-%m-%d"),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"requirements-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/docx")
async def export_requirements_docx(
    q: str | None = Query(None),
    status: list[str] | None = Query(None),
    priority: list[str] | None = Query(None),
    source: list[str] | None = Query(None),
    system_id: list[int] | None = Query(None),
    stakeholder_id: list[int] | None = Query(None),
    tag: list[str] | None = Query(None),
    confidence: list[str] | None = Query(None),
    db: AsyncSession = Depends(get_db),
):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError:
        raise HTTPException(status_code=503, detail="python-docx not installed")

    items, _ = await svc.list_requirements(
        db, q=q, status=status, priority=priority, source=source,
        system_id=system_id, stakeholder_id=stakeholder_id,
        tag=tag, confidence=confidence, page=1, page_size=10_000,
    )

    doc = Document()
    doc.add_heading("Requirements Report", 0)
    doc.add_paragraph(f"Generated: {date.today().isoformat()} | Total: {len(items)}")
    doc.add_paragraph("")

    for r in items:
        h = doc.add_heading(f"{r.req_id}: {r.title}", level=1)
        h.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        fields = [
            ("Status", r.status), ("Priority", r.priority), ("Source", r.source),
            ("Confidence", r.confidence),
            ("Stakeholder", r.stakeholder.name if r.stakeholder else "—"),
            ("System", r.system.name if r.system else "—"),
            ("Tags", ", ".join(t.name for t in r.tags) or "—"),
        ]
        hdr = table.rows[0].cells
        hdr[0].text = "Field"
        hdr[1].text = "Value"
        for field, value in fields:
            row = table.add_row().cells
            row[0].text = field
            row[1].text = str(value)
        doc.add_paragraph("Description:")
        doc.add_paragraph(r.description or "").runs[0].font.size = Pt(10) if r.description else None
        if r.business_impact:
            doc.add_paragraph(f"Business Impact: {r.business_impact}")
        if r.technical_impact:
            doc.add_paragraph(f"Technical Impact: {r.technical_impact}")
        doc.add_paragraph("")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = f"requirements-{date.today().isoformat()}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/import", status_code=200)
async def import_requirements_csv(
    file: UploadFile,
    db: AsyncSession = Depends(get_db),
):
    if not file.filename or not file.filename.endswith((".csv", ".xlsx")):
        raise HTTPException(status_code=422, detail="Only .csv and .xlsx files are supported")

    content = await file.read()
    created_ids = []
    errors = []

    try:
        if file.filename.endswith(".xlsx"):
            try:
                import openpyxl
            except ImportError:
                raise HTTPException(status_code=503, detail="openpyxl not installed")
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return {"created": [], "errors": ["Empty file"]}
            headers = [str(h).strip().lower() if h else "" for h in rows[0]]
            data_rows = [dict(zip(headers, row)) for row in rows[1:]]
        else:
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            data_rows = [{k.strip().lower(): v for k, v in row.items()} for row in reader]
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Failed to parse file: {e}")

    from app.schemas.common import ConfidenceEnum, PriorityEnum, SourceEnum
    valid_priorities = {e.value.lower(): e.value for e in PriorityEnum}
    valid_sources = {e.value.lower(): e.value for e in SourceEnum}
    valid_confidences = {e.value.lower(): e.value for e in ConfidenceEnum}

    for i, row in enumerate(data_rows, start=2):
        title = (row.get("title") or "").strip()
        if not title:
            errors.append(f"Row {i}: missing title")
            continue
        try:
            req_data = RequirementCreate(
                title=title,
                description=(row.get("description") or "").strip() or title,
                source=valid_sources.get((row.get("source") or "").strip().lower(), "Stakeholder Interview"),
                priority=valid_priorities.get((row.get("priority") or "").strip().lower(), "Medium"),
                confidence=valid_confidences.get((row.get("confidence") or "").strip().lower(), "Medium"),
                business_impact=(row.get("business impact") or row.get("business_impact") or "").strip() or None,
                technical_impact=(row.get("technical impact") or row.get("technical_impact") or "").strip() or None,
                notes=(row.get("notes") or "").strip() or None,
                tag_names=[t.strip() for t in (row.get("tags") or "").split(",") if t.strip()],
            )
            req = await svc.create_requirement(db, req_data)
            created_ids.append(req.req_id)
        except Exception as e:
            errors.append(f"Row {i}: {e}")

    return {"created": created_ids, "errors": errors}


@router.get("/export/pdf")
async def export_requirements_pdf(
    q: str | None = Query(None),
    status: list[str] | None = Query(None),
    priority: list[str] | None = Query(None),
    source: list[str] | None = Query(None),
    system_id: list[int] | None = Query(None),
    stakeholder_id: list[int] | None = Query(None),
    tag: list[str] | None = Query(None),
    confidence: list[str] | None = Query(None),
    db: AsyncSession = Depends(get_db),
):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    except ImportError:
        raise HTTPException(status_code=503, detail="reportlab not installed")

    items, _ = await svc.list_requirements(
        db, q=q, status=status, priority=priority, source=source,
        system_id=system_id, stakeholder_id=stakeholder_id,
        tag=tag, confidence=confidence, page=1, page_size=10_000,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontSize=18, spaceAfter=6)
    h1_style = ParagraphStyle("H1", parent=styles["Heading2"], fontSize=12, textColor=colors.HexColor("#1e3a5f"), spaceBefore=14, spaceAfter=4)
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=9, spaceAfter=4)
    label_style = ParagraphStyle("Label", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#64748b"), spaceBefore=2)

    story = [
        Paragraph("Requirements Report", title_style),
        Paragraph(f"Generated: {date.today().isoformat()} &nbsp;·&nbsp; Total: {len(items)}", body_style),
        Spacer(1, 0.4*cm),
    ]

    for r in items:
        story.append(Paragraph(f"{r.req_id}: {r.title}", h1_style))
        meta_data = [
            ["Status", r.status, "Priority", r.priority],
            ["Source", r.source, "Confidence", r.confidence],
            ["Stakeholder", r.stakeholder.name if r.stakeholder else "—", "System", r.system.name if r.system else "—"],
        ]
        tbl = Table(meta_data, colWidths=[2.5*cm, 5*cm, 2.5*cm, 5*cm])
        tbl.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#64748b")),
            ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#64748b")),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(tbl)
        if r.description:
            story.append(Paragraph("<b>Description:</b>", label_style))
            story.append(Paragraph(r.description.replace("\n", "<br/>"), body_style))
        if r.business_impact:
            story.append(Paragraph(f"<b>Business Impact:</b> {r.business_impact}", body_style))
        if r.technical_impact:
            story.append(Paragraph(f"<b>Technical Impact:</b> {r.technical_impact}", body_style))
        if r.tags:
            story.append(Paragraph(f"<b>Tags:</b> {', '.join(t.name for t in r.tags)}", body_style))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0"), spaceAfter=6))

    doc.build(story)
    buf.seek(0)
    filename = f"requirements-{date.today().isoformat()}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/bulk-status", response_model=list[RequirementResponse])
async def bulk_transition_status(
    data: BulkStatusRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not data.req_ids:
        raise HTTPException(status_code=422, detail="No requirement IDs provided")
    try:
        new_status = StatusEnum(data.status)
    except ValueError:
        raise HTTPException(status_code=422, detail=f"Invalid status: {data.status}")

    results = []
    errors = []
    for req_id in data.req_ids:
        try:
            req = await svc.transition_status(db, req_id, new_status, changed_by_id=current_user.id)
            results.append(_to_response(req))
        except HTTPException as e:
            errors.append(f"{req_id}: {e.detail}")

    if errors and not results:
        raise HTTPException(status_code=422, detail="; ".join(errors))
    return results


@router.get("/{req_id}/versions", response_model=list[VersionResponse])
async def list_versions(req_id: str, db: AsyncSession = Depends(get_db)):
    return await ver_svc.list_versions(db, req_id)


@router.post("", response_model=RequirementResponse, status_code=201)
async def create_requirement(data: RequirementCreate, db: AsyncSession = Depends(get_db)):
    req = await svc.create_requirement(db, data)
    return _to_response(req)


@router.get("/{req_id}", response_model=RequirementResponse)
async def get_requirement(req_id: str, db: AsyncSession = Depends(get_db)):
    req = await svc.get_requirement(db, req_id)
    return _to_response(req)


@router.put("/{req_id}", response_model=RequirementResponse)
async def update_requirement(
    req_id: str,
    data: RequirementUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    req = await svc.update_requirement(db, req_id, data, changed_by_id=current_user.id)
    return _to_response(req)


@router.patch("/{req_id}/status", response_model=RequirementResponse)
async def transition_status(
    req_id: str,
    data: StatusTransition,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    req = await svc.transition_status(db, req_id, data.status, changed_by_id=current_user.id)
    return _to_response(req)


@router.delete("/{req_id}", status_code=204)
async def delete_requirement(req_id: str, db: AsyncSession = Depends(get_db)):
    await svc.delete_requirement(db, req_id)


@router.post("/{req_id}/relations/{target_req_id}", response_model=RequirementResponse)
async def add_relation(req_id: str, target_req_id: str, db: AsyncSession = Depends(get_db)):
    req = await svc.add_relation(db, req_id, target_req_id)
    return _to_response(req)


@router.delete("/{req_id}/relations/{target_req_id}", response_model=RequirementResponse)
async def remove_relation(req_id: str, target_req_id: str, db: AsyncSession = Depends(get_db)):
    req = await svc.remove_relation(db, req_id, target_req_id)
    return _to_response(req)


@router.post("/{req_id}/jira")
async def push_to_jira(req_id: str, db: AsyncSession = Depends(get_db)):
    req = await svc.get_requirement(db, req_id)
    try:
        return await jira_service.push_requirement(db, req)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@router.post("/{req_id}/linear")
async def push_to_linear(req_id: str, db: AsyncSession = Depends(get_db)):
    req = await svc.get_requirement(db, req_id)
    try:
        return await linear_service.push_requirement(db, req)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
